"""Excel rendering for collected candidate records."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Candidate, Language
from .text import clean_text


def format_positions(candidate: Candidate) -> str:
    if not candidate.positions:
        return candidate.experience_fallback
    blocks: list[str] = []
    for index, item in enumerate(candidate.positions, 1):
        lines = [f"{index}. {clean_text(item.title) or '未标注职位'}"]
        values = [
            ("公司", item.company),
            ("时间", item.date_range),
            ("任职时长", item.duration),
            ("地点", item.location),
            ("职责/成果", item.description),
            ("相关技能", "、".join(item.skills)),
        ]
        lines.extend(f"{label}：{clean_text(value)}" for label, value in values if clean_text(value))
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def format_education(candidate: Candidate) -> str:
    if not candidate.education:
        return candidate.education_fallback
    blocks: list[str] = []
    for index, item in enumerate(candidate.education, 1):
        lines = [f"{index}. {clean_text(item.school) or '未标注学校'}"]
        values = [
            ("学位", item.degree),
            ("专业", item.field_of_study),
            ("时间", item.date_range),
            ("活动", item.activities),
            ("成绩", item.grade),
        ]
        lines.extend(f"{label}：{clean_text(value)}" for label, value in values if clean_text(value))
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def format_languages(languages: list[Language]) -> str:
    return " | ".join(
        f"{item.name}（{item.proficiency}）" if item.proficiency else item.name
        for item in languages
        if item.name
    )


def excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return value[:32_767]


def export_excel(records: list[Candidate], output: Path) -> None:
    headers = [
        "姓名", "当前职位", "当前公司", "地点", "Headline",
        "LinkedIn Public URL", "Recruiter URL", "核心技能",
        "完整工作经历", "完整教育经历", "求职开放信息", "语言", "母语为中文",
    ]
    widths = [24, 32, 30, 24, 52, 48, 52, 44, 76, 62, 48, 36, 16]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "候选人"
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        row = [
            record.name,
            record.current_title,
            record.current_company,
            record.location,
            record.headline,
            record.public_url,
            record.recruiter_url,
            " | ".join(record.core_skills),
            format_positions(record),
            format_education(record),
            "\n".join(record.open_to_work),
            format_languages(record.languages),
            record.native_chinese,
        ]
        sheet.append([excel_safe(value) for value in row])
        row_number = sheet.max_row
        for column in range(1, len(headers) + 1):
            sheet.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=True)
        for column in (6, 7):
            cell = sheet.cell(row_number, column)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
